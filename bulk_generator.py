"""
bulk_generator.py
-----------------
Enterprise bulk offer-letter generation engine (no UI here).

Responsibilities
================
    * Import candidate rows from .xlsx / .csv with flexible column mapping.
    * Normalise + validate rows (dates, required fields).
    * Generate hundreds/thousands of PDFs reliably, in a cancellable run that
      reports progress through a callback (so the UI thread never blocks).
    * Write a date-organised folder tree and an audit log.

The engine is deliberately UI-agnostic and reuses the existing
:class:`pdf_generator.PDFGenerator`, company settings and templates, so it can
later be driven from a CLI, a scheduler or an email-sending pipeline without
changes.

Columns recognised (case-insensitive, extra columns ignored):
    Candidate Name | Position | Domain | Issue Date | Start Date | End Date
"""

from __future__ import annotations

import csv
import logging
import threading
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import utils
from pdf_generator import PDFGenerator
import doc_router

try:
    import openpyxl
except ImportError:  # xlsx support optional; csv still works.
    openpyxl = None  # type: ignore


# --------------------------------------------------------------------------- #
# Column model
# --------------------------------------------------------------------------- #
# Canonical field -> list of accepted header aliases (lower-cased, stripped).
COLUMN_ALIASES: Dict[str, List[str]] = {
    "intern_id": ["intern id", "intern_id", "internid", "id"],
    "candidate_name": ["candidate name", "name", "candidate", "full name"],
    "position": ["position", "role", "designation", "job title"],
    "domain": ["domain", "department", "dept", "team"],
    "issue_date": ["issue date", "issued", "date of issue", "letter date"],
    "start_date": ["start date", "start", "from", "joining date"],
    "end_date": ["end date", "end", "to", "completion date"],
}

DISPLAY_COLUMNS: List[Tuple[str, str]] = [
    ("candidate_name", "Candidate Name"),
    ("position", "Position"),
    ("issue_date", "Issue Date"),
    ("start_date", "Start Date"),
    ("end_date", "End Date"),
    ("intern_id", "Intern ID"),
    ("issues", "Issues"),
]

# Column headers written by :func:`write_sample` and understood by the importer.
SAMPLE_HEADERS: List[str] = ["Candidate Name", "Position", "Department",
                             "Issue Date", "Start Date", "End Date",
                             "Intern ID"]

# Columns the user must not hand-edit in the table (system allocated or derived).
READONLY_COLUMNS = {"intern_id", "issues"}


# --------------------------------------------------------------------------- #
# Logging
# --------------------------------------------------------------------------- #
def get_logger(base_dir: Path) -> logging.Logger:
    """Return a singleton file logger writing to ``logs/generation.log``."""
    log_dir = base_dir / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    logger = logging.getLogger("offerletter.bulk")
    if not logger.handlers:
        handler = logging.FileHandler(log_dir / "generation.log", encoding="utf-8")
        handler.setFormatter(logging.Formatter(
            "%(asctime)s | %(levelname)-7s | %(message)s",
            datefmt="%Y-%m-%d %H:%M:%S"))
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
    return logger


# --------------------------------------------------------------------------- #
# Import / normalisation
# --------------------------------------------------------------------------- #
def _norm_header(value: object) -> str:
    return str(value or "").strip().lower()


def _build_index(headers: List[object]) -> Dict[str, int]:
    """Map canonical field names to their column index in ``headers``."""
    index: Dict[str, int] = {}
    normalised = [_norm_header(h) for h in headers]
    for field_name, aliases in COLUMN_ALIASES.items():
        for i, header in enumerate(normalised):
            if header in aliases:
                index[field_name] = i
                break
    return index


def _norm_date(value: object) -> str:
    """Normalise any date-ish cell to ``DD-MM-YYYY`` (best effort)."""
    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime(utils.DATE_FORMAT)
    text = str(value).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d-%b-%Y",
                "%d %B %Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).strftime(utils.DATE_FORMAT)
        except ValueError:
            continue
    return text  # leave untouched; validation will flag it


# Public alias: other modules (intern_io) normalise dates the same way, and
# reaching for a private name across modules is worse than exporting one.
normalize_date = _norm_date


def _row_from(values: List[object], index: Dict[str, int]) -> Dict[str, str]:
    def cell(field_name: str) -> object:
        i = index.get(field_name)
        return values[i] if i is not None and i < len(values) else ""

    return {
        "candidate_name": str(cell("candidate_name") or "").strip(),
        "position": str(cell("position") or "").strip(),
        "domain": str(cell("domain") or "").strip(),
        "issue_date": _norm_date(cell("issue_date")),
        "start_date": _norm_date(cell("start_date")),
        "end_date": _norm_date(cell("end_date")),
        # Optional: when a sheet carries a known Intern ID we reuse it as-is
        # instead of allocating a new one.
        "intern_id": str(cell("intern_id") or "").strip(),
    }


def write_sample(path: str | Path) -> Path:
    """Write a starter spreadsheet with the exact headers the importer reads.

    Saves people guessing the column names - which is the most common reason an
    import comes back empty.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    today = datetime.now()
    example = ["Asha Verma", "AI Agent Developer Intern", "Engineering",
               utils.format_date(today), "01-07-2026", "30-09-2026", ""]

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import Workbook
        except ImportError as exc:                      # pragma: no cover
            raise ValueError("openpyxl is needed to write an .xlsx sample") \
                from exc
        book = Workbook()
        sheet = book.active
        sheet.title = "Candidates"
        sheet.append(SAMPLE_HEADERS)
        sheet.append(example)
        for i, header in enumerate(SAMPLE_HEADERS, start=1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=i).column_letter].width = \
                max(len(header) + 4, 14)
        book.save(path)
        return path

    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(SAMPLE_HEADERS)
        writer.writerow(example)
    return path


def zip_folder(folder: str | Path, archive: str | Path) -> Path:
    """Zip every file in ``folder`` (non-recursive) into ``archive``."""
    import zipfile

    folder = Path(folder)
    archive = Path(archive)
    archive.parent.mkdir(parents=True, exist_ok=True)
    files = sorted(f for f in folder.iterdir()
                   if f.is_file() and f.suffix.lower() != ".zip")
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in files:
            zf.write(file, arcname=file.name)
    return archive


def import_rows(path: str | Path) -> List[Dict[str, str]]:
    """Read candidate rows from an .xlsx or .csv file.

    Raises ``ValueError`` for unsupported extensions or a missing engine.
    """
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        return _import_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _import_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix} (use .xlsx or .csv)")


def _import_csv(path: Path) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        all_rows = [r for r in reader if any(str(c).strip() for c in r)]
    if not all_rows:
        return rows
    index = _build_index(all_rows[0])
    for values in all_rows[1:]:
        row = _row_from(values, index)
        if any(row.values()):
            rows.append(row)
    return rows


def _import_xlsx(path: Path) -> List[Dict[str, str]]:
    if openpyxl is None:
        raise ValueError("openpyxl is required for .xlsx import. "
                         "Install it with: pip install openpyxl")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows: List[Dict[str, str]] = []
    header: Optional[List[object]] = None
    index: Dict[str, int] = {}
    for raw in ws.iter_rows(values_only=True):
        values = list(raw)
        if header is None:
            if any(str(c).strip() for c in values if c is not None):
                header = values
                index = _build_index(header)
            continue
        if not any(str(c).strip() for c in values if c is not None):
            continue
        row = _row_from(values, index)
        if any(row.values()):
            rows.append(row)
    wb.close()
    return rows


# --------------------------------------------------------------------------- #
# Validation
# --------------------------------------------------------------------------- #
def validate_row(row: Dict[str, str]) -> List[str]:
    """Return a list of human readable problems for a single row (empty = ok)."""
    errors: List[str] = []
    if not row.get("candidate_name", "").strip():
        errors.append("Missing Name")
    if not row.get("position", "").strip():
        errors.append("Missing Position")

    issue = row.get("issue_date", "").strip()
    if not issue:
        errors.append("Missing Issue Date")
    elif utils.parse_date(issue) is None:
        errors.append("Invalid Issue Date")

    start = row.get("start_date", "").strip()
    if not start:
        errors.append("Empty Start Date")
    elif utils.parse_date(start) is None:
        errors.append("Invalid Start Date")

    end = row.get("end_date", "").strip()
    if not end:
        errors.append("Missing End Date")
    elif utils.parse_date(end) is None:
        errors.append("Invalid End Date")

    s, e = utils.parse_date(start), utils.parse_date(end)
    if s and e and e < s:
        errors.append("End before Start")
    return errors


def validate_all(rows: List[Dict[str, str]]) -> Dict[int, List[str]]:
    """Return ``{row_index: errors}`` for every invalid row."""
    return {i: errs for i, row in enumerate(rows)
            if (errs := validate_row(row))}


# --------------------------------------------------------------------------- #
# Folder organisation
# --------------------------------------------------------------------------- #
def output_dir_for(base_output: Path, template_label: str,
                   when: Optional[datetime] = None) -> Path:
    """Create and return ``output/<YYYY-MM-DD>/<Category>/`` for a template."""
    when = when or datetime.now()
    date_folder = when.strftime("%Y-%m-%d")
    category = ("Certificates" if "Certificate" in template_label
                else "Offer Letters")
    folder = base_output / date_folder / category
    folder.mkdir(parents=True, exist_ok=True)
    return folder


# --------------------------------------------------------------------------- #
# Generation engine
# --------------------------------------------------------------------------- #
@dataclass
class BulkReport:
    """Outcome of a bulk run, also used to export a CSV report."""
    total: int = 0
    succeeded: List[Dict[str, str]] = field(default_factory=list)
    failed: List[Dict[str, str]] = field(default_factory=list)
    cancelled: bool = False
    seconds: float = 0.0
    folder: str = ""

    @property
    def success_count(self) -> int:
        return len(self.succeeded)

    @property
    def fail_count(self) -> int:
        return len(self.failed)

    def export_csv(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(["Candidate Name", "Intern ID", "Certificate No",
                             "File", "Status", "Detail", "Time"])
            for r in self.succeeded:
                writer.writerow([r["candidate_name"], r.get("intern_id", ""),
                                 r.get("cert_no", ""), r["file"], "Generated",
                                 "", r["time"]])
            for r in self.failed:
                writer.writerow([r["candidate_name"], r.get("intern_id", ""),
                                 r.get("cert_no", ""), r.get("file", ""),
                                 "Failed", r.get("error", ""), r["time"]])


# Progress callback signature: (done:int, total:int, current_name:str)
ProgressCb = Callable[[int, int, str], None]


class BulkGenerator:
    """Generates many PDFs reusing the shared PDF engine and company settings."""

    def __init__(self, base_dir: Path, output_base: Path, db=None):
        self.base_dir = base_dir
        self.output_base = output_base
        self.pdf = PDFGenerator(base_dir)
        self.logger = get_logger(base_dir)
        self.db = db  # optional InternDatabase for ID allocation + records
        # Certificate store (shares the same DB file) for completion certs.
        self.cert_store = None
        if db is not None:
            try:
                from certificate_generator import CertificateStore
                self.cert_store = CertificateStore(db.db_path)
            except Exception:
                self.cert_store = None

    def _to_form_data(self, row: Dict[str, str]) -> Dict[str, str]:
        """Map an import row to the data dict the templates expect."""
        duration = utils.duration_between(row.get("start_date", ""),
                                          row.get("end_date", ""))
        return {
            "candidate_name": row.get("candidate_name", ""),
            "position": row.get("position", ""),
            "department": row.get("domain", ""),
            "course": "",
            "college": "",
            "issue_date": row.get("issue_date", ""),
            "start_date": row.get("start_date", ""),
            "end_date": row.get("end_date", ""),
            "duration": duration,
            "email": "",
            # Pre-known Intern ID (rows loaded from the database carry theirs).
            "intern_id": str(row.get("intern_id", "") or "").strip(),
        }

    def generate(self, rows: List[Dict[str, str]], company: Dict,
                 template_label: str, progress_cb: Optional[ProgressCb] = None,
                 cancel_event: Optional[threading.Event] = None) -> BulkReport:
        """Generate a PDF for every row. Safe to call from a worker thread."""
        report = BulkReport(total=len(rows))
        started = datetime.now()
        folder = output_dir_for(self.output_base, template_label, started)
        report.folder = str(folder)
        self.logger.info("=== Bulk run started: %d rows, template=%s ===",
                         len(rows), template_label)

        for i, row in enumerate(rows):
            if cancel_event is not None and cancel_event.is_set():
                report.cancelled = True
                self.logger.warning("Bulk run cancelled at %d/%d", i, len(rows))
                break

            name = row.get("candidate_name", "Candidate")
            stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                data = self._to_form_data(row)
                data["template"] = template_label
                parsed = utils.parse_date(data.get("issue_date", ""))
                year = parsed.year if parsed else datetime.now().year
                # Intern ID: trust the one already on the row (candidates
                # picked from the database), else reuse by name, else allocate.
                if self.db is not None and not data.get("intern_id"):
                    existing = self.db.get_by_name(name)
                    if existing and existing.get("intern_id"):
                        data["intern_id"] = existing["intern_id"]
                    else:
                        data["intern_id"] = self.db.next_intern_id(year)
                # Every certificate type gets a verifiable certificate number,
                # drawn from its own family (SO-INT / SO-CERT).
                if doc_router.is_certificate(template_label) and \
                        self.cert_store is not None:
                    data["cert_no"] = self.cert_store.next_cert_no(
                        year, doc_router.cert_kind(template_label))
                base = doc_router.filename_base(template_label, name)
                # unique_path guarantees (1), (2)… - never overwrites.
                path = utils.unique_path(folder, base)
                doc_router.render(path, company, data, template_label,
                                  self.base_dir, self.pdf)
                if self.db is not None:
                    self.db.add_record({
                        "intern_id": data.get("intern_id", ""),
                        "candidate_name": name,
                        "position": data.get("position", ""),
                        "domain": data.get("department", ""),
                        "issue_date": data.get("issue_date", ""),
                        "start_date": data.get("start_date", ""),
                        "end_date": data.get("end_date", ""),
                        "duration": data.get("duration", ""),
                        "pdf_path": str(path),
                        "status": "Generated",
                    })
                if doc_router.is_certificate(template_label) and \
                        self.cert_store is not None:
                    ongoing = doc_router.is_ongoing_cert(template_label)
                    self.cert_store.add_record({
                        **data, "pdf_path": str(path),
                        "cert_type": template_label,
                        "status": ("Internship Ongoing" if ongoing
                                   else "Completed")})
                report.succeeded.append(
                    {"candidate_name": name, "file": path.name, "time": stamp,
                     "intern_id": data.get("intern_id", ""),
                     "cert_no": data.get("cert_no", "")})
                self.logger.info("OK   | %-30s | %s", name, path.name)
            except Exception as exc:  # never let one bad row stop the batch
                report.failed.append(
                    {"candidate_name": name, "error": str(exc), "time": stamp})
                self.logger.error("FAIL | %-30s | %s", name, exc)

            if progress_cb is not None:
                progress_cb(i + 1, len(rows), name)

        report.seconds = (datetime.now() - started).total_seconds()
        self.logger.info("=== Bulk run finished: %d ok, %d failed, %.1fs ===",
                         report.success_count, report.fail_count, report.seconds)
        return report
