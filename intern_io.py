"""
intern_io.py
------------
Import and export of intern records **together with their certificates**.

Why a separate module
=====================
``bulk_generator`` reads spreadsheets of *candidates to generate documents
for*. This module moves *records that already exist* in and out of the
database - a different job with different rules, so it lives on its own and
never changes how bulk generation behaves.

The database holds two tables in one file (``company/interns.db``):

    interns       keyed on ``intern_id``
    certificates  keyed on ``cert_no``, pointing back at an ``intern_id``

An export therefore has to carry both, and an import has to put both back
without breaking the ID allocators or silently overwriting existing work.

Formats
=======
================  ==========================================================
``.xlsx``         Two sheets, ``Interns`` and ``Certificates``. Full
                  round-trip. The recommended format.
``.json``         ``{"interns": [...], "certificates": [...]}``. Full
                  round-trip, exact values, easy to diff or script.
``.csv``          Interns only - a single CSV cannot hold two tables. The
                  caller is told plainly rather than losing data quietly.
================  ==========================================================

Import is deliberately a two-step operation: :func:`preview` reports what
*would* happen, the caller confirms, then :func:`apply_archive` writes. Both
``add_record`` methods are ``INSERT OR REPLACE``, so overwriting existing
records is possible and must never happen without the user agreeing to it.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import bulk_generator as bg
import utils

try:
    import openpyxl
except ImportError:                                   # xlsx optional
    openpyxl = None                                   # type: ignore


APP_LABEL = "CertiFlow"
SCHEMA_VERSION = 1

# --------------------------------------------------------------------------- #
# Column model
# --------------------------------------------------------------------------- #
# (field, human header) in the order they are written. The field names match
# the database columns exactly, so a record maps straight through.
INTERN_COLUMNS: List[Tuple[str, str]] = [
    ("intern_id", "Intern ID"),
    ("candidate_name", "Candidate Name"),
    ("position", "Position"),
    ("domain", "Department"),
    ("start_date", "Start Date"),
    ("end_date", "End Date"),
    ("issue_date", "Issue Date"),
    ("duration", "Duration"),
    ("status", "Status"),
    ("created_at", "First Added"),
    ("pdf_path", "Document File"),
]

CERT_COLUMNS: List[Tuple[str, str]] = [
    ("cert_no", "Certificate No"),
    ("intern_id", "Intern ID"),
    ("candidate_name", "Candidate Name"),
    ("cert_type", "Certificate Type"),
    ("position", "Position"),
    ("program", "Program"),
    ("department", "Department"),
    ("start_date", "Start Date"),
    ("end_date", "End Date"),
    ("issue_date", "Issue Date"),
    ("duration", "Duration"),
    ("grade", "Grade"),
    ("remarks", "Remarks"),
    ("status", "Status"),
    ("created_at", "Issued At"),
    ("pdf_path", "Certificate File"),
]

INTERN_SHEET = "Interns"
CERT_SHEET = "Certificates"

# Header aliases so a hand-made or re-saved sheet still imports. Extends the
# bulk importer's aliases with the record-only fields.
_EXTRA_ALIASES: Dict[str, List[str]] = {
    "cert_no": ["certificate no", "certificate number", "cert no", "cert_no",
                "certno", "certificate id"],
    "cert_type": ["certificate type", "cert type", "cert_type", "type",
                  "document type"],
    "program": ["program", "programme", "program name"],
    "department": ["department", "dept", "domain", "team"],
    "duration": ["duration", "period", "length"],
    "grade": ["grade", "rating"],
    "remarks": ["remarks", "remark", "comments", "note", "notes"],
    "status": ["status", "state"],
    "created_at": ["issued at", "first added", "created at", "created_at",
                   "recorded at"],
    "pdf_path": ["certificate file", "document file", "pdf", "pdf path",
                 "pdf_path", "file"],
}

# Fields that carry a date and should be normalised to DD-MM-YYYY on import.
_DATE_FIELDS = ("start_date", "end_date", "issue_date")

# The minimum needed for a row to be worth importing.
_REQUIRED_INTERN = ("intern_id", "candidate_name")
_REQUIRED_CERT = ("cert_no", "candidate_name")


def _aliases_for(fields: List[str]) -> Dict[str, List[str]]:
    """Alias table covering ``fields``, preferring the bulk importer's list."""
    table: Dict[str, List[str]] = {}
    for name in fields:
        aliases = list(bg.COLUMN_ALIASES.get(name, []))
        aliases += [a for a in _EXTRA_ALIASES.get(name, []) if a not in aliases]
        # The canonical header itself always matches.
        for source in (INTERN_COLUMNS, CERT_COLUMNS):
            for f, header in source:
                if f == name and header.lower() not in aliases:
                    aliases.append(header.lower())
        table[name] = aliases
    return table


_INTERN_ALIASES = _aliases_for([f for f, _ in INTERN_COLUMNS])
_CERT_ALIASES = _aliases_for([f for f, _ in CERT_COLUMNS])


# --------------------------------------------------------------------------- #
# Archive
# --------------------------------------------------------------------------- #
@dataclass
class Archive:
    """Records read from, or destined for, a file."""
    interns: List[Dict[str, str]] = field(default_factory=list)
    certificates: List[Dict[str, str]] = field(default_factory=list)
    source: str = ""
    # Rows dropped while reading, as ``(what, why)`` for reporting.
    skipped: List[Tuple[str, str]] = field(default_factory=list)

    def __bool__(self) -> bool:
        return bool(self.interns or self.certificates)


@dataclass
class Preview:
    """What an import would do, before anything is written."""
    new_interns: List[Dict[str, str]] = field(default_factory=list)
    existing_interns: List[Dict[str, str]] = field(default_factory=list)
    new_certs: List[Dict[str, str]] = field(default_factory=list)
    existing_certs: List[Dict[str, str]] = field(default_factory=list)
    # Certificates whose Intern ID matches no intern, in the file or on record.
    orphan_certs: List[Dict[str, str]] = field(default_factory=list)
    skipped: List[Tuple[str, str]] = field(default_factory=list)

    @property
    def total_new(self) -> int:
        return len(self.new_interns) + len(self.new_certs)

    @property
    def total_existing(self) -> int:
        return len(self.existing_interns) + len(self.existing_certs)


@dataclass
class Result:
    """What an import actually did."""
    interns_added: int = 0
    interns_updated: int = 0
    certs_added: int = 0
    certs_updated: int = 0
    skipped: List[Tuple[str, str]] = field(default_factory=list)
    failed: List[Tuple[str, str]] = field(default_factory=list)

    @property
    def total(self) -> int:
        return (self.interns_added + self.interns_updated
                + self.certs_added + self.certs_updated)


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def _clean(value: Any) -> str:
    """A cell as trimmed text; ``None`` becomes empty."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime(utils.DATE_FORMAT)
    return str(value).strip()


def _index_for(headers: List[Any], aliases: Dict[str, List[str]]
               ) -> Dict[str, int]:
    """Map field names to column positions using ``aliases``."""
    normalised = [_clean(h).lower() for h in headers]
    index: Dict[str, int] = {}
    for name, names in aliases.items():
        for i, header in enumerate(normalised):
            if header and header in names:
                index[name] = i
                break
    return index


def _record_from(values: List[Any], index: Dict[str, int],
                 columns: List[Tuple[str, str]]) -> Dict[str, str]:
    record: Dict[str, str] = {}
    for name, _header in columns:
        position = index.get(name)
        raw = (values[position]
               if position is not None and position < len(values) else "")
        record[name] = (bg.normalize_date(raw) if name in _DATE_FIELDS
                        else _clean(raw))
    return record


def _looks_like_certificates(headers: List[Any]) -> bool:
    """True when a sheet's headers include a certificate number column."""
    return "cert_no" in _index_for(headers, {"cert_no": _CERT_ALIASES["cert_no"]})


def _rows_of(sheet) -> List[List[Any]]:
    """Non-blank rows of a worksheet."""
    rows = []
    for raw in sheet.iter_rows(values_only=True):
        values = list(raw)
        if any(_clean(c) for c in values):
            rows.append(values)
    return rows


def _validate(record: Dict[str, str], required: Tuple[str, ...],
              label: str) -> Optional[str]:
    """Return a reason to skip ``record``, or ``None`` when it is usable."""
    missing = [name for name in required if not record.get(name, "").strip()]
    if missing:
        pretty = " and ".join(
            dict(INTERN_COLUMNS + CERT_COLUMNS).get(m, m) for m in missing)
        return f"{label}: missing {pretty}"
    return None


def _collect(rows: List[List[Any]], columns: List[Tuple[str, str]],
             aliases: Dict[str, List[str]], required: Tuple[str, ...],
             label: str) -> Tuple[List[Dict[str, str]], List[Tuple[str, str]]]:
    """Turn header + data rows into validated records."""
    if not rows:
        return [], []
    index = _index_for(rows[0], aliases)
    if not index:
        return [], [(label, "no recognised column headers")]
    records: List[Dict[str, str]] = []
    skipped: List[Tuple[str, str]] = []
    for values in rows[1:]:
        record = _record_from(values, index, columns)
        if not any(record.values()):
            continue
        problem = _validate(record, required, label)
        if problem:
            who = (record.get("candidate_name") or record.get("intern_id")
                   or record.get("cert_no") or "row")
            skipped.append((who, problem))
            continue
        records.append(record)
    return records, skipped


def read_archive(path: str | Path) -> Archive:
    """Read interns and certificates from ``path``.

    Raises ``ValueError`` for an unsupported extension or a missing engine.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".json":
        return _read_json(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv(path)
    raise ValueError(
        f"Unsupported file type: {suffix or path.name} "
        f"(use .xlsx, .csv or .json)")


def _read_json(path: Path) -> Archive:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"That file is not valid JSON ({exc.msg}).") from exc
    if not isinstance(payload, dict):
        raise ValueError("Expected a JSON object with 'interns' and "
                         "'certificates'.")

    archive = Archive(source=path.name)
    intern_fields = [f for f, _ in INTERN_COLUMNS]
    cert_fields = [f for f, _ in CERT_COLUMNS]

    for raw in payload.get("interns") or []:
        if not isinstance(raw, dict):
            continue
        record = {f: _clean(raw.get(f, "")) for f in intern_fields}
        problem = _validate(record, _REQUIRED_INTERN, "Intern")
        if problem:
            archive.skipped.append((record.get("candidate_name") or "record",
                                    problem))
            continue
        archive.interns.append(record)

    for raw in payload.get("certificates") or []:
        if not isinstance(raw, dict):
            continue
        record = {f: _clean(raw.get(f, "")) for f in cert_fields}
        problem = _validate(record, _REQUIRED_CERT, "Certificate")
        if problem:
            archive.skipped.append((record.get("cert_no") or "record", problem))
            continue
        archive.certificates.append(record)
    return archive


def _read_xlsx(path: Path) -> Archive:
    if openpyxl is None:
        raise ValueError("openpyxl is required to read .xlsx files. "
                         "Install it with: pip install openpyxl")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    archive = Archive(source=path.name)
    try:
        for sheet in book.worksheets:
            rows = _rows_of(sheet)
            if not rows:
                continue
            # Decide by the headers, not the sheet name, so a renamed or
            # re-saved workbook still imports correctly.
            if _looks_like_certificates(rows[0]):
                records, skipped = _collect(rows, CERT_COLUMNS, _CERT_ALIASES,
                                            _REQUIRED_CERT, "Certificate")
                archive.certificates += records
            else:
                records, skipped = _collect(rows, INTERN_COLUMNS,
                                            _INTERN_ALIASES,
                                            _REQUIRED_INTERN, "Intern")
                archive.interns += records
            archive.skipped += skipped
    finally:
        book.close()
    return archive


def _read_csv(path: Path) -> Archive:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = [r for r in csv.reader(fh) if any(_clean(c) for c in r)]
    archive = Archive(source=path.name)
    if not rows:
        return archive
    if _looks_like_certificates(rows[0]):
        archive.certificates, archive.skipped = _collect(
            rows, CERT_COLUMNS, _CERT_ALIASES, _REQUIRED_CERT, "Certificate")
    else:
        archive.interns, archive.skipped = _collect(
            rows, INTERN_COLUMNS, _INTERN_ALIASES, _REQUIRED_INTERN, "Intern")
    return archive


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #
def csv_holds_interns_only(path: str | Path) -> bool:
    """True when writing to ``path`` would drop the certificates.

    Lets the UI warn *before* the user picks a format that cannot carry
    everything, instead of losing records silently.
    """
    return Path(path).suffix.lower() == ".csv"


def _autosize(sheet, headers: List[str]) -> None:
    for i, header in enumerate(headers, start=1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=i).column_letter].width = \
            max(len(header) + 4, 14)


def write_archive(path: str | Path, interns: List[Dict[str, str]],
                  certificates: List[Dict[str, str]]) -> Path:
    """Write ``interns`` and ``certificates`` to ``path``.

    ``.csv`` writes the interns only - a flat file cannot represent two
    tables. Use :func:`csv_holds_interns_only` to warn first.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix.lower()

    if suffix == ".json":
        payload = {
            "app": APP_LABEL,
            "schema": SCHEMA_VERSION,
            "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "interns": [{f: r.get(f, "") for f, _ in INTERN_COLUMNS}
                        for r in interns],
            "certificates": [{f: r.get(f, "") for f, _ in CERT_COLUMNS}
                             for r in certificates],
        }
        path.write_text(json.dumps(payload, indent=2, ensure_ascii=False),
                        encoding="utf-8")
        return path

    if suffix in (".xlsx", ".xlsm"):
        if openpyxl is None:
            raise ValueError("openpyxl is required to write .xlsx files. "
                             "Install it with: pip install openpyxl")
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = INTERN_SHEET
        headers = [h for _f, h in INTERN_COLUMNS]
        sheet.append(headers)
        for record in interns:
            sheet.append([_clean(record.get(f, "")) for f, _ in INTERN_COLUMNS])
        _autosize(sheet, headers)

        certs_sheet = book.create_sheet(CERT_SHEET)
        cert_headers = [h for _f, h in CERT_COLUMNS]
        certs_sheet.append(cert_headers)
        for record in certificates:
            certs_sheet.append([_clean(record.get(f, ""))
                                for f, _ in CERT_COLUMNS])
        _autosize(certs_sheet, cert_headers)

        # Freeze the header row on both sheets - these files get scrolled.
        sheet.freeze_panes = "A2"
        certs_sheet.freeze_panes = "A2"
        book.save(path)
        return path

    if suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow([h for _f, h in INTERN_COLUMNS])
            for record in interns:
                writer.writerow([_clean(record.get(f, ""))
                                 for f, _ in INTERN_COLUMNS])
        return path

    raise ValueError(f"Unsupported file type: {suffix or path.name} "
                     f"(use .xlsx, .csv or .json)")


def export_from_db(path: str | Path, interns: List[Dict]) -> Tuple[Path, int, int]:
    """Write ``interns`` (as returned by ``intern_directory.load_interns``).

    Returns ``(path, intern_count, certificate_count)``. Each intern already
    carries its ``certificates`` list, so both tables come from one read.
    """
    flat_interns = [{f: record.get(f, "") or "" for f, _ in INTERN_COLUMNS}
                    for record in interns]
    certificates: List[Dict[str, str]] = []
    seen = set()
    for record in interns:
        for cert in record.get("certificates") or []:
            key = (cert.get("cert_no") or "").strip().upper()
            # A certificate re-attached by name can appear under two interns.
            if key and key in seen:
                continue
            if key:
                seen.add(key)
            certificates.append({f: cert.get(f, "") or ""
                                 for f, _ in CERT_COLUMNS})
    written = write_archive(path, flat_interns, certificates)
    return written, len(flat_interns), len(certificates)


def write_sample(path: str | Path) -> Path:
    """A starter file showing the exact columns the importer understands."""
    today = utils.format_date(datetime.now())
    intern = {
        "intern_id": "SO260001", "candidate_name": "Asha Verma",
        "position": "AI Agent Developer Intern", "domain": "Engineering",
        "start_date": "01-07-2026", "end_date": "30-09-2026",
        "issue_date": today, "duration": "3 Months",
        "status": "Generated", "created_at": "", "pdf_path": "",
    }
    cert = {
        "cert_no": "SO-INT-260001", "intern_id": "SO260001",
        "candidate_name": "Asha Verma", "cert_type": "Internship Certificate",
        "position": "AI Agent Developer Intern", "program": "",
        "department": "Engineering", "start_date": "01-07-2026",
        "end_date": "30-09-2026", "issue_date": today,
        "duration": "3 Months", "grade": "", "remarks": "",
        "status": "Internship Ongoing", "created_at": "", "pdf_path": "",
    }
    return write_archive(path, [intern], [cert])


# --------------------------------------------------------------------------- #
# Importing
# --------------------------------------------------------------------------- #
def preview(db, cert_store, archive: Archive) -> Preview:
    """Classify every record in ``archive`` against what is already stored.

    Reads only - nothing is written. This is what lets the UI say exactly how
    many records would be *overwritten* before the user commits.
    """
    result = Preview(skipped=list(archive.skipped))

    known_interns = {(r.get("intern_id") or "").strip().upper()
                     for r in db.search("")}
    known_certs = {(r.get("cert_no") or "").strip().upper()
                   for r in cert_store.all_records()}

    # Later duplicates within the file itself are dropped, first one wins.
    seen_interns = set()
    for record in archive.interns:
        key = record["intern_id"].strip().upper()
        if key in seen_interns:
            result.skipped.append((record.get("candidate_name") or key,
                                   f"Intern: {record['intern_id']} appears "
                                   f"more than once in the file"))
            continue
        seen_interns.add(key)
        if key in known_interns:
            result.existing_interns.append(record)
        else:
            result.new_interns.append(record)

    # Certificates may attach to an intern already stored or one in this file.
    available = known_interns | seen_interns
    seen_certs = set()
    for record in archive.certificates:
        key = record["cert_no"].strip().upper()
        if key in seen_certs:
            result.skipped.append((record.get("cert_no") or key,
                                   f"Certificate: {record['cert_no']} appears "
                                   f"more than once in the file"))
            continue
        seen_certs.add(key)
        intern_key = (record.get("intern_id") or "").strip().upper()
        if not intern_key or intern_key not in available:
            result.orphan_certs.append(record)
        if key in known_certs:
            result.existing_certs.append(record)
        else:
            result.new_certs.append(record)
    return result


def apply_archive(db, cert_store, archive: Archive,
                  overwrite: bool = False) -> Result:
    """Write ``archive`` into the database.

    ``overwrite`` decides what happens to records that already exist: when
    ``False`` they are left exactly as they are and counted as skipped.

    Every imported Intern ID is reserved with the allocator, so an ID that came
    in from a file can never later be handed to a different person.
    """
    plan = preview(db, cert_store, archive)
    result = Result(skipped=list(plan.skipped))

    def store_intern(record: Dict[str, str]) -> bool:
        payload = {f: record.get(f, "") for f, _ in INTERN_COLUMNS}
        db.add_record(payload)
        # Keep the ID allocator ahead of anything imported.
        try:
            db.reserve_counter(payload.get("intern_id", ""))
        except Exception:
            pass
        return True

    def store_cert(record: Dict[str, str]) -> bool:
        payload = {f: record.get(f, "") for f, _ in CERT_COLUMNS}
        cert_store.add_record(payload)
        reserve = getattr(cert_store, "reserve_counter", None)
        if callable(reserve):
            try:
                reserve(payload.get("cert_no", ""))
            except Exception:
                pass
        return True

    for record in plan.new_interns:
        try:
            store_intern(record)
            result.interns_added += 1
        except Exception as exc:
            result.failed.append((record.get("intern_id", "?"), str(exc)))

    for record in plan.existing_interns:
        if not overwrite:
            result.skipped.append(
                (record.get("candidate_name") or record.get("intern_id", "?"),
                 f"Intern: {record.get('intern_id','')} already on record"))
            continue
        try:
            store_intern(record)
            result.interns_updated += 1
        except Exception as exc:
            result.failed.append((record.get("intern_id", "?"), str(exc)))

    for record in plan.new_certs:
        try:
            store_cert(record)
            result.certs_added += 1
        except Exception as exc:
            result.failed.append((record.get("cert_no", "?"), str(exc)))

    for record in plan.existing_certs:
        if not overwrite:
            result.skipped.append(
                (record.get("cert_no", "?"),
                 f"Certificate: {record.get('cert_no','')} already on record"))
            continue
        try:
            store_cert(record)
            result.certs_updated += 1
        except Exception as exc:
            result.failed.append((record.get("cert_no", "?"), str(exc)))

    return result
